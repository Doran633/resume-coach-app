"""Current compilation and persisted rendering controls, not historical replay."""
from copy import deepcopy
from dataclasses import replace
import json
from pathlib import Path

import pytest
from docx import Document

from app import models, schemas
from app.services import docx_service
from app.services.canonical_semantic_state_service import build_canonical_semantic_build
from app.services.canonical_consumer_view_service import build_canonical_consumer_views
from app.services.input_claim_resolution_service import resolve_experience_claims
from app.services.prompt_service import _canonical_evidence_context
from test_v09131_immutable_delivery_rendering_consistency import _payload, _database
from test_v09162_model_output_evidence_contract import request, CASES
from test_v09174_fact_reference_composition import isolated
from test_v09172_initial_evidence_preservation import detail_return
from test_v09162_model_output_evidence_contract import reference_return
from test_v09176_delivery_closure import deliver, provider
from test_v0916_canonical_model_evidence import evidence
from test_v091851_semantic_unit_boundary import RAW_INPUT
from test_v09184_denied_claim_validation import HISTORICAL_INPUT, bound, evaluate, denied
from app.services import generation_service as generation


@pytest.mark.parametrize('text', [
    '研究尚未形成论文，我也没有提出新的模型结构',
    '项目尚未上线，我也没有负责部署',
    '研究目前尚未形成论文',
])
def test_independent_denials_are_constraints(text, isolated):
    raw = '科研经历：声音分类课题\n使用Python整理音频数据。' + text + '。'
    build = build_canonical_semantic_build(raw)
    before = deepcopy(build)
    views = build_canonical_consumer_views(build)
    claims = [c for owner in views.experience_ids for c in views.claims_for_owner(owner)]
    denied = [c for c in claims if '尚未' in c.text or '也没有' in c.text]
    assert denied and all(c.eligibility == 'excluded' and c.polarity == 'negative' for c in denied)
    assert all(raw[slice(*c.source_span)] == c.text for c in denied)
    assert all('尚未' not in f.resume_ready_text and '也没有' not in f.resume_ready_text for f in build.ledger.facts)
    sent = _canonical_evidence_context(request(dict(CASES[0], raw_input=raw)), views)
    evidence = json.JSONDecoder().raw_decode(sent.split('\n', 1)[1])[0]
    assert {c.claim_id for c in denied} <= {c['claim_id'] for c in evidence['internal_constraints_not_resume_facts']}
    assert build == before


@pytest.mark.parametrize('text', [
    '使用Python研究尚未分类的音频',
    '实现没有网络时的本地查询',
    '工具仅在本地运行',
    '仅供同学试用',
    '研究形成了实验记录',
])
def test_actions_objects_and_project_status_remain_facts(text):
    result = resolve_experience_claims('EXP-001', text)
    assert result.claims and all(c.eligibility == 'eligible' for c in result.claims)


def test_persisted_docx_uses_neutral_rows_without_changing_payload(tmp_path, monkeypatch, isolated):
    payload = _payload()
    before = payload.model_dump()
    db = _database(payload)
    stored = db.get(models.GenerationResult, 9131).result_json
    monkeypatch.setattr(docx_service, 'OUTPUT_DIR', tmp_path)
    response = docx_service.create_docx(db, schemas.DocxCreate(
        anonymous_user_id='u', session_id='rendering-v09131', generation_result_id=9131))
    rows = [p.text for p in Document(tmp_path / response.file_name).paragraphs]
    assert '完成用户调研与数据复盘。' in rows
    assert '负责访谈记录整理。' in rows
    assert '整理用户反馈并形成分析结论。' in rows
    assert '技术细节：保留实际正文。' in rows
    assert not any(p.startswith(('经历简介：', '项目简介：', '我的职责：')) for p in rows)
    assert db.get(models.GenerationResult, 9131).result_json == stored
    assert payload.model_dump() == before
    db.close()


@pytest.mark.parametrize('separator', ['。', '。\n', '。\n\n', '。\r\n'])
def test_independent_positive_and_denials_keep_local_spans(separator):
    text = separator.join(['使用Python整理音频', '研究尚未形成论文，我也没有提出新的模型结构', '工具仅在本地运行'])
    result = resolve_experience_claims('EXP-001', text)
    assert [c.text for c in result.eligible_claims] == ['使用Python整理音频', '工具仅在本地运行']
    assert len(result.excluded_claims) == 2
    assert all(text[slice(*c.source_span)] == c.text for c in result.claims)
    assert result == resolve_experience_claims('EXP-001', text)


def test_positive_action_before_also_denied_is_not_lost():
    text = '完成实验记录，我也没有负责部署'
    result = resolve_experience_claims('EXP-001', text)
    assert [c.text for c in result.eligible_claims] == ['完成实验记录']
    assert [c.text for c in result.excluded_claims] == ['我也没有负责部署']
    assert all(text[slice(*c.source_span)] == c.text for c in result.claims)


@pytest.mark.parametrize('raw', [RAW_INPUT, HISTORICAL_INPUT], ids=['java', 'research'])
def test_full_existing_inputs_save_every_fact_and_export(monkeypatch, tmp_path, isolated, raw):
    case, build, data = detail_return(raw)
    data['resume_sections']['summary'] = ['具备项目开发实践。']
    for key in ('normal_version', 'bold_version', 'boundary_version', 'recommended_version'):
        data[key] = ''
    before = deepcopy(build)
    outcome = deliver(monkeypatch, tmp_path, case, [(json.dumps(reference_return(data), ensure_ascii=False), 'stop')])
    assert outcome['results'] == 1 and 'docx' in outcome, outcome
    assert build == before
    projects = outcome['saved']['resume_sections']['projects']
    for fact in build.ledger.facts:
        project = next(p for p in projects if p['source_experience_id'] == fact.experience_id)
        assert fact.fact_id in project['source_fact_ids']
        assert fact.resume_ready_text.strip(' 。') in json.dumps(project, ensure_ascii=False)
    rows = [p.text for p in Document(next(tmp_path.glob('*.docx'))).paragraphs]
    assert not any(p.startswith(('项目简介：', '经历简介：', '我的职责：', '技术细节：')) for p in rows)


@pytest.mark.parametrize('long_mode', [False, True])
def test_retry_uses_same_full_evidence_and_rejects_omission(monkeypatch, isolated, long_mode):
    case, build, data = detail_return(RAW_INPUT)
    good = reference_return(data)
    bad = deepcopy(good)
    bad['resume_sections']['projects'][0]['fact_placements'].popitem()
    sent = provider(monkeypatch, [(json.dumps(bad, ensure_ascii=False), 'stop'), (json.dumps(good, ensure_ascii=False), 'stop')])
    views = build_canonical_consumer_views(build)
    before = deepcopy(build), repr(views)
    _, stats = generation.build_llm_generation(request(case), replace(build.long_input_context, long_input_mode=long_mode), consumer_views=views)
    assert stats['attempt'] == len(sent) == 2
    assert evidence(sent[0]['messages'][1]['content']) == evidence(sent[1]['messages'][1]['content'])
    assert (build, repr(views)) == before


def test_non_displayed_denial_still_rejects_contradiction(isolated):
    build, payload = bound('科研经历：声音分类课题\n使用Python整理音频。研究尚未形成论文。')
    project = payload.resume_sections.projects[0]
    assert not any('尚未' in line for line in project['details'])
    project['details'].append('研究形成论文')
    project['detail_fact_ids'].append([])
    project['detail_claim_ids'].append([])
    assert denied(evaluate(build, payload))


def test_web_neutral_preview_keeps_existing_scope():
    root = Path(__file__).resolve().parents[1]
    source = (root / 'frontend/src/pages/ResultPage.tsx').read_text(encoding='utf-8')
    preview = source.split('function ProjectPreview', 1)[1].split('function FactStrip', 1)[0]
    assert 'project.details.slice(0, 3)' in preview
    assert 'result.resume_sections.projects?.[0]' in source
    assert '<li key={index}>{item}</li>' in preview
    assert 'cleanDisplayText(project.intro)' not in preview
    assert 'raw_input' not in preview


def test_holdout_positive_work_and_independent_negation(monkeypatch, tmp_path, isolated):
    raw = ('项目经历：仓储记录核对工具\n2026年5月，使用Python核对仓储记录。'
           '使用17条记录检查重复编号，我也没有负责生产部署。'
           '工具仅在本地运行。')
    case, build, data = detail_return(raw)
    denied_claims = [c for c in build.ledger.claims if c.eligibility == 'excluded']
    assert any(c.text == '我也没有负责生产部署' for c in denied_claims)
    assert any('17条记录' in f.resume_ready_text for f in build.ledger.facts)
    data['resume_sections']['summary'] = ['具备记录核对工具开发实践。']
    for key in ('normal_version', 'bold_version', 'boundary_version', 'recommended_version'):
        data[key] = ''
    outcome = deliver(monkeypatch, tmp_path, case, [(json.dumps(reference_return(data), ensure_ascii=False), 'stop')])
    assert outcome['results'] == 1 and 'docx' in outcome, outcome
    project = outcome['saved']['resume_sections']['projects'][0]
    assert set(project['source_fact_ids']) == {f.fact_id for f in build.ledger.facts}
    assert '没有负责生产部署' not in json.dumps(project, ensure_ascii=False)
    assert '工具仅在本地运行' in json.dumps(project, ensure_ascii=False)

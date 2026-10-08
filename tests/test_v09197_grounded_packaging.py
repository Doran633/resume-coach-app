"""Actual network boundary controls; support verdicts do not prove model reliability."""
from copy import deepcopy
from dataclasses import replace
import json
from pathlib import Path

from docx import Document
import pytest

from app import schemas
from app.services import generation_service as generation, experience_slot_service as slots
from app.services.canonical_consumer_view_service import build_canonical_consumer_views
from app.services.resume_skill_evidence_aggregation_service import aggregate_skill_evidence_from_ledger, skill_evidence_display
from app.services.resume_skill_taxonomy_service import _compact_declarations
from app.services.resume_language_professionalization_service import professionalize_resume_language
from test_v09174_fact_reference_composition import isolated
from test_v09176_delivery_closure import deliver, provider
from test_v09162_model_output_evidence_contract import request
from test_v09193_evidence_bounded_composition import composition_sample, decision, review_reply

EVIDENCE = json.loads((Path(__file__).parent / 'fixtures/v09197_thin_candidate_evidence.json').read_text(encoding='utf-8'))
CHECKS = {name: True for name in (
    'source_meaning_preserved', 'owner_scope_preserved', 'qualification_preserved',
    'no_new_concrete_action', 'no_new_hard_claim',
)}


def thin_sample():
    case, build, body = composition_sample(EVIDENCE['raw_input'], combine=False)
    body['resume_sections']['projects'] = deepcopy(EVIDENCE['projects'])
    actual_ids = {f.fact_id for f in build.ledger.facts}
    assert actual_ids == {fid for p in body['resume_sections']['projects']
                          for u in p['expression_units'] for fid in u['fact_ids']}
    return case, build, body


@pytest.mark.parametrize('index', [0, 1])
def test_recorded_thin_candidate_rejection_not_reclassified(index, monkeypatch, tmp_path, isolated):
    case, build, body = thin_sample()
    before = deepcopy(build)
    outcome = deliver(monkeypatch, tmp_path, case, [
        (json.dumps(body, ensure_ascii=False), 'stop'),
        (json.dumps(EVIDENCE['reviews'][index], ensure_ascii=False), 'stop'),
    ])
    assert outcome['error'] == 'MODEL_EXPRESSION_REJECTED'
    assert outcome['results'] == 0 and not list(tmp_path.glob('*.docx'))
    assert len(outcome['sent']) == 2
    assert build == before


@pytest.mark.parametrize('position', ['intro', 'role', 'detail'])
def test_equivalent_local_role_and_editorial_observation_save(position, monkeypatch, tmp_path, isolated):
    case, build, body = composition_sample(EVIDENCE['raw_input'], combine=False)
    unit = body['resume_sections']['projects'][0]['expression_units'][1]
    unit.update(position=position, text='承担使用 Vue 开发活动列表与报名表单这一部分工作，按同学提供的接口展示提交结果。')
    fid = unit['fact_ids'][0]
    source = next(f for f in build.ledger.facts if f.fact_id == fid)
    verdict = decision('editorial_observe', fid, source.resume_ready_text,
                       '承担使用 Vue 开发活动列表与报名表单这一部分工作')
    verdict['checks'] = CHECKS.copy()
    reply = review_reply(build, body, {fid: verdict})
    before = deepcopy(build)
    outcome = deliver(monkeypatch, tmp_path, case, [
        (json.dumps(body, ensure_ascii=False), 'stop'), (json.dumps(reply, ensure_ascii=False), 'stop'),
    ])
    assert outcome['results'] == 1 and 'docx' in outcome, outcome
    text = '\n'.join(p.text for p in Document(next(tmp_path.glob('*.docx'))).paragraphs)
    assert unit['text'].rstrip('。') in text
    assert unit['text'].rstrip('。') in outcome['saved']['bold_version']
    project = outcome['saved']['resume_sections']['projects'][0]
    rows = list(zip(project['details'], project['detail_fact_ids'], project['detail_claim_ids']))
    rows += [(project.get(p, ''), project.get(p + '_source_fact_ids', []),
              project.get(p + '_source_claim_ids', [])) for p in ('intro', 'role')]
    _, facts, claims = next(row for row in rows if unit['text'].rstrip('。') in row[0])
    assert fid in facts and source.claim_id in claims
    assert build == before and len(outcome['sent']) == 2


@pytest.mark.parametrize('verdict', ['added_claim', 'changed_qualification', 'uncertain'])
def test_hard_claim_duty_upgrade_and_major_uncertainty_do_not_save(verdict, monkeypatch, tmp_path, isolated):
    case, build, body = composition_sample(EVIDENCE['raw_input'], combine=False)
    unit = body['resume_sections']['projects'][0]['expression_units'][1]
    unit['text'] = '独立负责整体报名系统开发并部署上线。'
    fid = unit['fact_ids'][0]
    source = next(f for f in build.ledger.facts if f.fact_id == fid)
    reply = review_reply(build, body, {fid: decision(
        verdict, None if verdict == 'added_claim' else fid,
        None if verdict == 'added_claim' else source.resume_ready_text, unit['text'],
    )})
    outcome = deliver(monkeypatch, tmp_path, case, [
        (json.dumps(body, ensure_ascii=False), 'stop'), (json.dumps(reply, ensure_ascii=False), 'stop'),
    ])
    expected = 'MODEL_EXPRESSION_REVIEW_UNCERTAIN' if verdict == 'uncertain' else 'MODEL_EXPRESSION_REJECTED'
    assert outcome['error'] == expected
    assert outcome['results'] == 0 and not list(tmp_path.glob('*.docx'))
    assert len(outcome['sent']) == 2


@pytest.mark.parametrize('long_mode', [False, True])
def test_actual_tasks_share_semantic_role_and_grounded_summary_contract(long_mode, monkeypatch, isolated):
    case, build, body = composition_sample(EVIDENCE['raw_input'], combine=False)
    views = build_canonical_consumer_views(build)
    before = deepcopy(build), repr(views)
    sent = provider(monkeypatch, [('not json', 'stop'), (json.dumps(body, ensure_ascii=False), 'stop')])
    generation.build_llm_generation(request(case), replace(build.long_input_context, long_input_mode=long_mode), consumer_views=views)
    for wire in sent:
        actual = wire['messages'][1]['content']
        assert '职责按实际参与关系、负责对象和范围判断' in actual
        assert 'summary 从已有行动证据提炼能力与用途' in actual
        assert '能力结论 + 行动证据 + 价值结果' not in actual
        assert '不能把课堂演示自动写成没有上线' in actual
        evidence = json.JSONDecoder().raw_decode(actual.split('<canonical_model_evidence>\n')[1])[0]
        assert {f['fact_id'] for o in evidence['owners'] for f in o['eligible_facts']} == {f.fact_id for f in build.ledger.facts}
        assert evidence['non_experience_context_not_project_facts']
    assert (build, repr(views)) == before


def test_summary_is_not_a_second_downstream_writer(monkeypatch, tmp_path, isolated):
    raw = '课程项目：活动登记页面\n2026年6月，使用Vue制作活动列表。帮忙测试Vue页面，检查空输入提示。'
    case, build, body = composition_sample(raw, combine=False)
    source = next(f for f in build.ledger.facts if '帮忙测试' in f.resume_ready_text)
    body['resume_sections']['summary'] = [source.resume_ready_text]
    changes = []
    actual = generation.professionalize_resume_language
    def observe(payload, **kwargs):
        result = actual(payload, **kwargs)
        changes.append((payload.resume_sections.summary[:], result.resume_sections.summary[:]))
        return result
    monkeypatch.setattr(generation, 'professionalize_resume_language', observe)
    outcome = deliver(monkeypatch, tmp_path, case, [(json.dumps(body, ensure_ascii=False), 'stop')])
    assert changes and all(before == after for before, after in changes), changes
    assert outcome['results'] == 1 and 'docx' in outcome, outcome
    assert outcome['saved']['resume_sections']['summary'] == [source.resume_ready_text.rstrip('。')]
    assert '问题定位' not in outcome['saved']['bold_version']


def test_skill_display_repetition_is_compacted_without_losing_degree(monkeypatch, tmp_path, isolated):
    raw = ('个人项目：记录页面\n2026年6月，使用Vue制作记录列表和添加表单。\n\n'
           '技能：\n能够使用Python和Java，了解基础Linux命令。接触过QuartzKit，仅用于课程练习。')
    case, build, body = composition_sample(raw, combine=False)
    before = deepcopy(build)
    outcome = deliver(monkeypatch, tmp_path, case, [(json.dumps(body, ensure_ascii=False), 'stop')])
    assert outcome['results'] == 1 and 'docx' in outcome, outcome
    skills = outcome['saved']['resume_sections']['skills']
    languages = next(s for s in skills if s.startswith('编程语言：'))
    assert languages.count('能够使用') == 1 and 'Python' in languages and 'Java' in languages
    assert any('了解基础Linux命令' in s for s in skills)
    assert any('接触过QuartzKit，仅用于课程练习' in s for s in skills)
    text = '\n'.join(p.text for p in Document(next(tmp_path.glob('*.docx'))).paragraphs)
    assert all(s in text and s in outcome['saved']['bold_version'] for s in skills)
    assert build == before


@pytest.mark.parametrize('declaration,expected', [
    ('能够使用Python和Java。', ['能够使用Python、Java']),
    ('能够使用Python。能够使用Java。', ['能够使用Python', '能够使用Java']),
    ('会使用Python，了解Java。', ['会使用Python', '了解Java']),
    ('接触过Python和Java，仅用于课程练习。', ['接触过Python，仅用于课程练习', '接触过Java，仅用于课程练习']),
    ('会使用Python和Java。了解Python。', ['会使用Python；了解Python', '会使用Java']),
])
def test_skill_compaction_requires_identical_source_and_qualification(declaration, expected):
    raw = '个人项目：登记页面\n2026年6月，使用Vue开发登记表单。\n\n技能：\n' + declaration
    _, build, _ = composition_sample(raw, combine=False)
    views = build_canonical_consumer_views(build)
    rows = aggregate_skill_evidence_from_ledger(build.ledger, non_experience_context=views.non_experience_context(raw))
    languages = [row for row in rows if row.term in ('Python', 'Java')]
    before = deepcopy(rows), deepcopy(build), repr(views)
    assert _compact_declarations(languages) == expected
    assert _compact_declarations(languages) == expected
    assert (rows, build, repr(views)) == before


def test_summary_legacy_derivation_is_still_available():
    from test_resume_language_professionalization import payload as make_payload
    payload = make_payload([])
    payload.resume_sections.summary = ['帮忙测试Vue页面']
    legacy = professionalize_resume_language(payload, write_log=False)
    canonical = professionalize_resume_language(payload, write_log=False, preserve_summary_expressions=True)
    assert '问题定位' in legacy.resume_sections.summary[0]
    assert canonical.resume_sections.summary == payload.resume_sections.summary
